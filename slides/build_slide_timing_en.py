# -*- coding: utf-8 -*-
"""3. JIRA Registration Timing — English, Accenture (ACN) style.

Usage:
    python build_slide_timing_en.py [-o out.pptx]

Layout follows the Accenture 2020 grid measured earlier in this project:
0.42" side margins over a 12.50" text column, an Arial Black headline set flush
top-left with no rule under it, and a "So What" band closing the page.

Built on a bare page rather than the starter-pack master: the master's purple ">"
is Accenture's own mark, so it is not reproduced here. Supply the .potx and use
build_slide.py --brand acn if the real master furniture is wanted.
"""
import argparse

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---- Accenture 2020 palette -------------------------------------------------
BLACK, WHITE = C("000000"), C("FFFFFF")
INK = C("333333")
KEY = C("A100FF")        # accent1 - Accenture Purple
MID = C("7500C0")        # accent2
DARK = C("460073")       # accent3
LIGHT = C("BE82FF")      # accent5
TINT = C("F4E9FF")
WARM = C("E6E6DC")       # lt2
WARM_GREY = C("96968C")  # dk2
MUTED = C("70706A")

LAT, DISPLAY = "Arial", "Arial Black"

# ---- ACN grid ---------------------------------------------------------------
L, R = 0.42, 12.92
W = R - L
TRACK_L = 1.78           # lane tracks start after the row labels
N = 9
PITCH = (R - TRACK_L) / N
CHEV_W = PITCH - 0.045
bx = lambda i: TRACK_L + PITCH * i
cx = lambda i: TRACK_L + PITCH * i + PITCH / 2
SOP_X = bx(6)            # boundary between Validation and Mass Production
SPEC_X = cx(2)           # Spec FIX      -> PGG / MGG trigger
PLAN_X = cx(3)           # Development Plan -> INF trigger
AXIS_Y = 3.92

PHASES = ["Inquiry", "Proposal", "Spec\nReview", "Investment", "Prod.\nPrep",
          "Validation", "Mass\nProd.", "Service\nParts", "EOP"]


# ---- helpers ----------------------------------------------------------------
def style(run, size, bold=False, color=BLACK, face=LAT):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def textbox(sl, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tb, tf


def say(sl, x, y, w, h, text, size, bold=False, color=BLACK,
        align=PP_ALIGN.LEFT, face=LAT, anchor=MSO_ANCHOR.TOP, space=None):
    _, tf = textbox(sl, x, y, w, h, align, anchor)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face)
    return tf


def line(sl, x1, y1, x2, y2, color, width=1.0, dash=None):
    ln = sl.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if dash:
        el = ln.line._get_or_add_ln()
        el.append(el.makeelement(qn("a:prstDash"), {"val": dash}))
    return ln


def shape(sl, kind, x, y, w, h, fill=None, outline=None, ow=1.0):
    sh = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return sh


# ---- the slide --------------------------------------------------------------
def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])      # blank

    say(sl, L, 0.36, 9.5, 0.24,
        "3. JIRA Registration Timing | Trigger by case type", 11.5, True, KEY)
    say(sl, L, 0.62, 12.50, 0.66,
        "Register PGG/MGG at Spec FIX and INF at Development Plan issuance",
        20, True, BLACK, face=DISPLAY)
    say(sl, L, 1.34, 12.0, 0.28,
        "The registration trigger differs by case type; SOP marks the switch "
        "from JIRA tracking to per-case PCR.", 13, False, MUTED)

    # --- trigger definition band ---
    shape(sl, MSO_SHAPE.RECTANGLE, L, 1.78, W, 0.70, fill=TINT)
    _, tf = textbox(sl, L + 0.22, 1.90, 11.9, 0.24)
    p = tf.paragraphs[0]
    for t, sz, b, c in [("Trigger", 12, True, DARK), ("     ", 12, False, BLACK),
                        ("PGG / MGG = at Spec FIX", 12.5, True, BLACK),
                        ("     |     ", 12.5, False, WARM_GREY),
                        ("INF = at Development Plan issuance", 12.5, True, BLACK),
                        ("     |     ", 12.5, False, WARM_GREY),
                        ("Managed in JIRA from registration onward", 12.5, True, MID)]:
        r = p.add_run(); r.text = t; style(r, sz, b, c)
    say(sl, L + 0.22, 2.18, 11.9, 0.22,
        "Before registration each case is managed individually. After registration, "
        "progress and issues are tracked in JIRA.", 10.5, False, MUTED)

    # --- lifecycle ---
    say(sl, L, 2.64, 4.0, 0.24, "Project lifecycle", 11.5, True, KEY)
    for txt, x0, x1, col in [("Pre-SOP", bx(4), bx(6), MID),
                             ("SOP", SOP_X - 0.34, SOP_X + 0.34, DARK),
                             ("Post-SOP", bx(6), bx(8), DARK)]:
        say(sl, x0, 2.64, x1 - x0, 0.24, txt, 10.5, True, col, align=PP_ALIGN.CENTER)
    line(sl, SOP_X, 2.92, SOP_X, 5.44, DARK, 1.25, dash="dash")

    for i, name in enumerate(PHASES):
        if i in (0, 1, 8):
            fill, fg, out = WARM, MUTED, None
        elif i in (6, 7):
            fill, fg, out = (KEY if i == 6 else MID), WHITE, None
        else:
            fill, fg, out = TINT, DARK, KEY
        ch = shape(sl, MSO_SHAPE.CHEVRON, bx(i), 3.00, CHEV_W, 0.52,
                   fill=fill, outline=out, ow=1.0)
        for j, ln in enumerate(name.split("\n")):
            para = ch.text_frame.paragraphs[0] if j == 0 else ch.text_frame.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            para.line_spacing = 0.92
            r = para.add_run(); r.text = ln
            style(r, 8.5, True, fg)

    # --- milestone markers ---
    for x, label in [(SPEC_X, "Spec FIX"), (PLAN_X, "Development Plan")]:
        say(sl, x - 1.15, 3.58, 2.30, 0.22, label, 10.5, True, DARK,
            align=PP_ALIGN.CENTER)
    line(sl, TRACK_L, AXIS_Y, R, AXIS_Y, WARM_GREY, 0.75)
    for x in (SPEC_X, PLAN_X):
        shape(sl, MSO_SHAPE.DIAMOND, x - 0.075, AXIS_Y - 0.075, 0.15, 0.15, fill=KEY)

    # --- swim lanes ---
    for name, y, trig in [("PGG / MGG", 4.64, SPEC_X), ("INF", 5.26, PLAN_X)]:
        say(sl, L, y - 0.13, 1.28, 0.26, name, 13, True, BLACK,
            align=PP_ALIGN.RIGHT, face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        line(sl, TRACK_L, y, trig, y, WARM_GREY, 1.25)
        line(sl, trig, y, R, y, KEY, 1.75)
        line(sl, trig, AXIS_Y, trig, y, LIGHT, 0.75, dash="sysDot")
        shape(sl, MSO_SHAPE.DIAMOND, trig - 0.075, y - 0.075, 0.15, 0.15, fill=KEY)
        for txt, x0, x1, sz, b, col in [
                ("Managed individually", TRACK_L + 0.06, trig, 10, False, MUTED),
                ("Managed in JIRA", trig, SOP_X, 11, True, MID),
                ("PCR raised per case", SOP_X, R, 11, True, DARK)]:
            say(sl, x0, y - 0.34, x1 - x0, 0.22, txt, sz, b, col, align=PP_ALIGN.CENTER)

    # --- open items ---
    shape(sl, MSO_SHAPE.RECTANGLE, L, 5.62, W, 0.60, fill=WARM)
    say(sl, L + 0.22, 5.74, 1.5, 0.22, "Open items", 11.5, True, DARK)
    for txt, x0 in [("1  Rare cases in the registration decision", 1.95),
                    ("2  Specs can still change after Spec FIX", 5.60),
                    ("3  Treatment of new devices [TBC]", 9.25)]:
        say(sl, x0, 5.74, 3.55, 0.40, txt, 10.5, False, INK)

    # --- So What ---
    shape(sl, MSO_SHAPE.RECTANGLE, L, 6.32, W, 0.62, fill=KEY)
    say(sl, L + 0.30, 6.44, 1.6, 0.26, "So What", 13, True, WHITE,
        face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L + 2.05, 6.44, 10.3, 0.26,
        "Fix the two registration triggers and split management responsibility at SOP",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, 7.10, 6.0, 0.22,
        "JIRA Registration Timing | Trigger definition", 9, False, MUTED)
    say(sl, 12.30, 7.10, 0.62, 0.22, "8", 9, False, MUTED, align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "PGG/MGG cases are registered in JIRA at Spec FIX; INF cases at Development "
        "Plan issuance. Before that point each case is managed individually. "
        "From SOP onward, Mass Production and Service Parts are handled by raising "
        "a PCR per case. Open: rare registration cases, spec changes after Spec FIX, "
        "and the treatment of new devices.")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="jira_timing_en_acn.pptx")
    build(ap.parse_args().output)
