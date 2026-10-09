# -*- coding: utf-8 -*-
"""必要アカウント数の確認依頼と、Jira Standardの参考価格（2ページ）。

Usage:
    python build_deck_account_cost.py [-o out.pptx]

P1: 各拠点への記入依頼（現在／今後／想定時期）
P2: 公式月額価格と累進計算の考え方、150名の計算例
"""
import argparse
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
GREEN = C("00843D")
GREEN_PALE = C("E7F4EC")
AMBER = C("C55A11")
AMBER_FILL = C("FDF1E3")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
FAINT = C("F4F5F3")
WHITE = C("FFFFFF")

JP = "Meiryo"
L, R = 0.42, 12.92
W = R - L

EM = 0.158
LH = 0.20


def style(run, size, bold=False, color=INK, face=JP, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, spacing=spacing)
    return tf


def rich(sl, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=1.25):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = space
    for text, size, bold, color in parts:
        r = p.add_run(); r.text = text
        style(r, size, bold, color)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chip(sl, x, y, w, h, text, fill, size=11, fg=WHITE, outline=None, ow=1.0):
    sh = rect(sl, x, y, w, h, fill, outline, ow)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    style(r, size, True, fg)
    return sh


def lines(text, inner_w):
    return max(1, math.ceil(len(text) / max(6, int(inner_w / EM))))


def bullet(sl, x, y, w, text, marker=BLUE, color=INK):
    h = lines(text, w - 0.20) * LH
    rect(sl, x, y + 0.075, 0.07, 0.07, marker)
    say(sl, x + 0.20, y, w - 0.20, h, text, 11, False, color, space=1.25)
    return h


def header(sl, eyebrow, headline, sub):
    say(sl, L, 0.36, 10.5, 0.26, eyebrow, 11.5, True, BLUE, spacing=1.0)
    say(sl, L, 0.62, W, 0.56, headline, 20, True, INK, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L, 1.20, 12.3, 0.26, sub, 12.5, False, GREY)


def foot(sl, y, label, page):
    say(sl, L, y, 7.5, 0.24, label, 11, False, GREY)
    say(sl, 12.20, y, 0.72, 0.24, page, 11, False, GREY, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- P1
SITE_COLS = [("拠点", 2.20), ("現在の有償アカウント数", 3.30),
             ("今後の必要アカウント数", 3.30), ("人数見込みの想定時期", 3.70)]
SITES = [("HQ", "14", "要確認", "要確認"),
         ("DSSE", "要確認", "要確認", "要確認"),
         ("DSST", "0", "要確認", "要確認"),
         ("DSSC", "0", "要確認", "要確認"),
         ("DSSA", "0", "要確認", "要確認")]
NOTES = [
    "※ 今後の必要数は追加人数ではなく、各拠点で必要となる総数を記入してください。"
    "新規拠点は展開時期、既存拠点は人数変更の想定時期を記載してください。",
    "※ 複数拠点で同一アカウントを利用する場合は、重複人数を別途お知らせください。",
]
ASKS = [("回答期限", "［要確認：各拠点の回答期限］"),
        ("提出先", "［要確認：HQの取りまとめ担当］")]

# ---------------------------------------------------------------- P2
PRICE_COLS = [("人数帯", 4.16), ("現行 月額単価", 4.16), ("改定後 月額単価", 4.18)]
PRICE_ROWS = [("1〜100", "9.05", "9.70"),
              ("101〜250", "7.65", "8.20"),
              ("251〜1,000", "6.40", "6.85")]
TERMS = ["改定価格は2026年10月13日（太平洋時間）から適用",
         "1,001名以上の単価は公式価格表を参照",
         "費用試算上、5拠点を1契約に集約すると仮定",
         "最終費用は代理店見積・契約条件で確認"]


def page1(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "① 依頼 ｜ 必要アカウント数の確認",
           "5拠点展開に必要なアカウント数の確認をお願いします",
           "HQでは、各拠点の現在の有償アカウント数と今後の必要アカウント数を集約し、"
           "5拠点展開に伴う費用を試算する。")

    rect(sl, L, 1.56, W, 0.40, PALE)
    chip(sl, L, 1.56, 1.42, 0.40, "現在の状況", DARK, 11)
    say(sl, L + 1.58, 1.56, W - 1.72, 0.40,
        "現在はHQ・DSSEで利用しており、今後はDSST・DSSC・DSSAへの展開を予定している。",
        12, True, DARK, anchor=MSO_ANCHOR.MIDDLE)

    rich(sl, L, 2.12, 12.3, 0.24,
         [("各拠点への依頼｜", 11.5, True, DARK),
          ("以下の表に、現在の有償アカウント数、今後の必要アカウント数、"
           "その人数を想定する時期を記入し、HQへご連絡ください。", 11.5, False,
           INK)])

    ty, hh, rh = 2.48, 0.38, 0.42
    x = L
    for name, cwid in SITE_COLS:
        chip(sl, x, ty, cwid - 0.02, hh, name, DARK, 11)
        x += cwid
    y = ty + hh
    for row in SITES:
        x = L
        for j, (val, (_, cwid)) in enumerate(zip(row, SITE_COLS)):
            tbd = val == "要確認"
            rect(sl, x, y, cwid - 0.02, rh,
                 AMBER_FILL if tbd else (PALE if j == 0 else WHITE), RULE, 0.75)
            say(sl, x, y, cwid - 0.02, rh, val, 11, j == 0 or not tbd,
                AMBER if tbd else (DARK if j == 0 else INK),
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            x += cwid
        y += rh

    y += 0.10
    for n in NOTES:
        h = lines(n, W) * LH
        say(sl, L, y, W, h, n, 11, False, GREY, space=1.25)
        y += h + 0.04

    ay, ah = 5.90, 0.62
    aw = (W - 0.24) / 2
    for i, (label, val) in enumerate(ASKS):
        ax = L + i * (aw + 0.24)
        rect(sl, ax, ay, aw, ah, AMBER_FILL, AMBER, 1.0)
        chip(sl, ax + 0.14, ay + 0.14, 1.34, ah - 0.28, label, AMBER, 11)
        say(sl, ax + 1.62, ay, aw - 1.76, ah, val, 12, True, AMBER,
            anchor=MSO_ANCHOR.MIDDLE)

    foot(sl, 7.10, "依頼 ｜ 必要アカウント数の確認", "1")

    sl.notes_slide.notes_text_frame.text = (
        "HQでは、各拠点の現在の有償アカウント数と今後の必要アカウント数を集約し、"
        "5拠点展開に伴う費用を試算します。現在はHQとDSSEで利用しており、今後は"
        "DSST・DSSC・DSSAへの展開を予定しています。表に、現在の有償アカウント数、"
        "今後の必要アカウント数、その人数を想定する時期を記入してHQへご連絡ください。"
        "今後の必要数は追加人数ではなく各拠点で必要となる総数です。新規拠点は展開時期、"
        "既存拠点は人数変更の想定時期をご記載ください。複数拠点で同一アカウントを"
        "利用する場合は、重複人数を別途お知らせください。")


def page2(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "② 参考価格 ｜ 月額費用の計算方法",
           "Jira Standardの参考価格と月額費用の計算方法",
           "各拠点の回答を基に、重複を除いた必要アカウント数で費用を試算する。")

    say(sl, L, 1.56, 8.0, 0.24, "公式月額価格（USD／人・月）", 11.5, True, DARK)
    ty, hh, rh = 1.86, 0.36, 0.40
    x = L
    for name, cwid in PRICE_COLS:
        chip(sl, x, ty, cwid - 0.02, hh, name, DARK, 11)
        x += cwid
    y = ty + hh
    for band, cur, new in PRICE_ROWS:
        x = L
        for j, val in enumerate((band, cur, new)):
            cwid = PRICE_COLS[j][1]
            rect(sl, x, y, cwid - 0.02, rh,
                 PALE if j == 2 else (FAINT if j == 0 else WHITE), RULE, 0.75)
            say(sl, x, y, cwid - 0.02, rh, val, 12 if j else 11.5, j != 1,
                DARK if j != 1 else INK, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
            x += cwid
        y += rh

    rich(sl, L, y + 0.10, W, 0.24,
         [("計算方法｜", 11, True, DARK),
          ("月額は人数帯ごとの累進計算。各人数帯に該当する人数に単価を掛け、"
           "合算する。", 11, False, GREY)])

    # 計算例
    ey, eh = 3.86, 1.24
    rect(sl, L, ey, W, eh, GREEN_PALE, GREEN, 1.0)
    chip(sl, L, ey, 2.70, 0.34, "計算例：150名の場合", GREEN, 11.5)
    bw, ow = 3.30, 0.62
    bx = L + (W - (bw * 2 + 3.70 + ow * 2)) / 2
    for i, (txt, wd) in enumerate(((("100名 × USD 9.70"), bw),
                                   ("＋", ow),
                                   ("50名 × USD 8.20", bw),
                                   ("＝", ow),
                                   ("改定後 USD 1,380／月", 3.70))):
        if txt in ("＋", "＝"):
            say(sl, bx, ey + 0.48, wd, 0.46, txt, 15, True, GREEN,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            last = i == 4
            rect(sl, bx, ey + 0.48, wd, 0.46, GREEN if last else WHITE,
                 GREEN, 1.0 if last else 0.75)
            say(sl, bx, ey + 0.48, wd, 0.46, txt, 13 if last else 12, True,
                WHITE if last else INK, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        bx += wd

    say(sl, L + 0.16, ey + eh - 0.26, W - 0.32, 0.22,
        "※ 150名は計算方法を示す例であり、5拠点の人数見込みではありません。",
        11, False, GREY)

    # 試算の前提
    say(sl, L, 5.32, 6.0, 0.24, "試算の前提", 11.5, True, DARK)
    rect(sl, L, 5.58, W, 0.02, BLUE)
    cw = (W - 0.30) / 2
    for i, t in enumerate(TERMS):
        tx = L + (i % 2) * (cw + 0.30)
        tyy = 5.70 + (i // 2) * 0.30
        bullet(sl, tx, tyy, cw, t)

    say(sl, L, 6.52, W, 0.24,
        "出典：Atlassian 公式価格表　"
        "https://www.atlassian.com/ja/licensing/future-pricing/cloud/list/"
        "pricing-tables", 11, False, GREY)

    foot(sl, 7.10, "参考価格 ｜ 月額費用の計算方法", "2")

    sl.notes_slide.notes_text_frame.text = (
        "各拠点の回答を基に、重複を除いた必要アカウント数で費用を試算します。"
        "公式の月額価格は1〜100名がUSD9.05から9.70へ、101〜250名が7.65から8.20へ、"
        "251〜1,000名が6.40から6.85へ改定されます。月額は人数帯ごとの累進計算で、"
        "各人数帯に該当する人数に単価を掛けて合算します。150名であれば、100名に"
        "9.70ドル、残る50名に8.20ドルを掛けて、改定後は月額1,380ドルとなります。"
        "この150名は計算方法を示す例であり、5拠点の人数見込みではありません。"
        "改定価格は2026年10月13日太平洋時間から適用、1,001名以上の単価は公式価格表を"
        "参照します。費用試算上は5拠点を1契約に集約すると仮定しており、最終費用は"
        "代理店見積と契約条件で確認します。")


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    page1(prs)
    page2(prs)
    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="account_cost_daicel.pptx")
    build(ap.parse_args().output)
