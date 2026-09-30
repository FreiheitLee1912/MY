# -*- coding: utf-8 -*-
"""会議体運営｜生準マネジメントサイクル（ACN版面 × ダイセル配色）

Usage:
    python build_slide_meeting_cadence.py [-o out.pptx]

第1週 Bi-weekly① → 第1〜2週 要因確認・分析 → 第3週 Bi-weekly② → 月末 生準月次報告
の1か月サイクルを4カラムで並べ、下に翌月へのループと運用の注記を置く。
"""
import argparse
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
MID = C("5B8CA8")          # 担当者作業（会議ではない工程）
AMBER = C("C55A11")
AMBER_FILL = C("FDF1E3")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
WHITE = C("FFFFFF")

JP = "Meiryo"
L, R = 0.42, 12.92
W = R - L

# 11pt Meiryo の全角1文字 ≒ 0.153"、行送り1.25 で 1行 ≒ 0.20"
EM11 = 0.153
LH11 = 0.20


def style(run, size, bold=False, color=INK, face=JP, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        face=JP, anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face, spacing=spacing)
    return tf


def rich(sl, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=1.25):
    """1段落に書体違いのrunを並べる（「目的｜…」のようなラベル付き行）。"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = space
    for text, size, bold, color in parts:
        r = p.add_run(); r.text = text
        style(r, size, bold, color)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chip(sl, x, y, w, h, text, fill, size=11, fg=WHITE):
    sh = rect(sl, x, y, w, h, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    style(r, size, True, fg)
    return sh


def lines(text, inner_w):
    """11pt 本文が inner_w に何行で収まるかを全角換算で見積もる。"""
    cpl = max(6, int(inner_w / EM11))
    return max(1, math.ceil(len(text) / cpl))


def bullet(sl, x, y, w, text, color=INK, marker=BLUE):
    h = lines(text, w - 0.20) * LH11
    rect(sl, x, y + 0.075, 0.07, 0.07, marker)
    say(sl, x + 0.20, y, w - 0.20, h, text, 11, False, color, space=1.25)
    return h


COLS = [
    dict(week="第1週", tag="会議", tagfill=DARK, hdr=DARK, w=2.76,
         label="Bi-weekly ①", name="変化点確認・問題提起",
         aim="前回以降の変化を把握する",
         todo=["カテゴリ別に案件の変化点を確認（日程・仕様・体制・コストなど）",
               "変化によって出た遅れ・問題を抜き出す",
               "前月の月次指摘のフォロー",
               "次回までの確認事項と担当を決める"],
         out="変化点と要フォロー案件"),
    dict(week="第1〜2週", tag="担当者作業", tagfill=MID, hdr=MID, w=2.76,
         label=None, name="要因確認・分析",
         aim="影響と打ち手を整理する",
         todo=["変化点の要因と影響範囲を確認",
               "リカバリ策と見通しを検討",
               "他案件・他カテゴリへの波及を確認"],
         out="影響とリカバリ案"),
    dict(week="第3週", tag="会議", tagfill=DARK, hdr=DARK, w=2.76,
         label="Bi-weekly ②", name="分析報告・まとめ",
         aim="方針の合意と月次の準備",
         todo=["①以降の新しい変化点を確認",
               "影響・リカバリ策を報告し方針を合意",
               "マネジメントへの依頼・確認事項を整理",
               "月次報告資料をまとめる"],
         out="月次報告資料"),
    dict(week="月末（第4週）", tag="報告", tagfill=AMBER, hdr=DARK, w=3.74,
         label="生準月次", name="報告",
         aim="マネジメントへの報告と意思決定",
         todo=["カテゴリ別に今月の主な変化点を報告",
               "影響の大きい案件と対応状況を報告",
               "決定・指摘事項を確認"],
         box=("マネジメントへの依頼・確認事項",
              ["判断してほしい事項（方針・優先順位）",
               "支援の依頼（人・予算・他部署との調整）",
               "前回依頼事項への回答確認"]),
         out="決定事項・指摘事項・依頼への回答"),
]

NOTES = [
    ("変化点を中心に確認", "変化のない案件は「変化なし」で済ませる"),
    ("変化点の観点", "日程・仕様・体制・コスト・外部要因"),
    ("月次に上げる基準", "影響の大きい変化点と、依頼・確認事項に絞る"),
]

GAP = 0.16
WEEK_Y, WEEK_H = 1.98, 0.30
CARD_Y, CARD_H = 2.32, 3.30
OUT_Y, OUT_H = 5.66, 0.48
LOOP_Y, LOOP_H = 6.20, 0.36
NOTE_Y, NOTE_H = 6.64, 0.44


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # --- ヘッダー ---
    say(sl, L, 0.36, 9.5, 0.26, "② 会議体運営 ｜ 生準マネジメントサイクル",
        11.5, True, BLUE, spacing=1.0)
    say(sl, L, 0.62, W, 0.56,
        "変化点を起点に、Bi-weekly ①②・要因分析・月次報告を1か月で一巡させる",
        20, True, INK, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L, 1.20, 12.0, 0.26,
        "各会議の議題は「前回からの変化」に絞り、月次の決定・指摘を翌月の入口に戻す。",
        12.5, False, GREY)

    # --- 基本方針バンド ---
    rect(sl, L, 1.54, W, 0.38, PALE)
    chip(sl, L, 1.54, 0.98, 0.38, "基本方針", DARK, 11.5)
    say(sl, L + 1.16, 1.54, W - 1.30, 0.38,
        "各会議では、前回からの「案件の変化点」を中心に確認する",
        12.5, True, DARK, anchor=MSO_ANCHOR.MIDDLE)

    # --- 4カラム ---
    xs, x = [], L
    for c in COLS:
        xs.append(x)
        x += c["w"] + GAP

    for i, (c, cx) in enumerate(zip(COLS, xs)):
        cw = c["w"]

        # 週ラベル帯＋区分タグ
        rect(sl, cx, WEEK_Y, cw, WEEK_H, PALE)
        say(sl, cx + 0.14, WEEK_Y, cw - 1.20, WEEK_H, c["week"], 11.5, True,
            DARK, anchor=MSO_ANCHOR.MIDDLE)
        tw = 0.18 + len(c["tag"]) * EM11
        chip(sl, cx + cw - 0.10 - tw, WEEK_Y + 0.04, tw, WEEK_H - 0.08,
             c["tag"], c["tagfill"], 11)

        # 流れの矢印（カラム間）
        if i < len(COLS) - 1:
            tri = rect(sl, cx + cw + 0.01, WEEK_Y + 0.06, 0.14, 0.18, BLUE,
                       shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            tri.rotation = 90

        # カード本体
        rect(sl, cx, CARD_Y, cw, CARD_H, WHITE, RULE, 0.75)

        # ヘッダー
        rect(sl, cx, CARD_Y, cw, 0.52, c["hdr"])
        if c["label"]:
            say(sl, cx + 0.14, CARD_Y + 0.05, cw - 0.28, 0.20, c["label"],
                11, True, PALE)
            say(sl, cx + 0.14, CARD_Y + 0.25, cw - 0.28, 0.24, c["name"],
                13, True, WHITE)
        else:
            say(sl, cx + 0.14, CARD_Y, cw - 0.28, 0.52, c["name"], 13, True,
                WHITE, anchor=MSO_ANCHOR.MIDDLE)

        # 目的
        rect(sl, cx + 0.01, CARD_Y + 0.52, cw - 0.02, 0.44, PALE)
        rich(sl, cx + 0.14, CARD_Y + 0.52, cw - 0.28, 0.44,
             [("目的｜", 11, True, DARK), (c["aim"], 11, False, INK)],
             anchor=MSO_ANCHOR.MIDDLE)

        # やること
        by = CARD_Y + 0.52 + 0.44 + 0.12
        say(sl, cx + 0.14, by, cw - 0.28, 0.20, "やること", 11, True, BLUE)
        by += 0.26
        for j, t in enumerate(c["todo"]):
            by += bullet(sl, cx + 0.14, by, cw - 0.28, t)
            if j < len(c["todo"]) - 1:
                by += 0.06

        # マネジメントへの依頼・確認事項（月末カラムのみ）
        if c.get("box"):
            title, items = c["box"]
            ih = [lines(t, cw - 0.56) * LH11 for t in items]
            bh = 0.08 + 0.20 + 0.06 + sum(ih) + 0.05 * (len(items) - 1) + 0.08
            bx, byy = cx + 0.12, by + 0.10
            rect(sl, bx, byy, cw - 0.24, bh, AMBER_FILL, AMBER, 0.75)
            say(sl, bx + 0.12, byy + 0.08, cw - 0.48, 0.20, title, 11, True,
                AMBER)
            iy = byy + 0.08 + 0.20 + 0.06
            for t, h in zip(items, ih):
                rect(sl, bx + 0.14, iy + 0.075, 0.07, 0.07, AMBER)
                say(sl, bx + 0.34, iy, cw - 0.58, h, t, 11, False, INK,
                    space=1.25)
                iy += h + 0.05

        # アウトプット
        rect(sl, cx, OUT_Y, cw, OUT_H, PALE, RULE, 0.75)
        rect(sl, cx, OUT_Y, 0.07, OUT_H, BLUE)
        rich(sl, cx + 0.20, OUT_Y, cw - 0.34, OUT_H,
             [("アウトプット｜", 11, True, DARK), (c["out"], 11, True, INK)],
             anchor=MSO_ANCHOR.MIDDLE, space=1.05)

    # --- 翌月へのループ ---
    rect(sl, L, LOOP_Y, W, LOOP_H, BLUE)
    chip(sl, L + 0.14, LOOP_Y + 0.05, 1.18, LOOP_H - 0.10, "サイクル", DARK, 11)
    say(sl, L + 1.50, LOOP_Y, W - 1.70, LOOP_H,
        "月次の決定・指摘事項と依頼への回答は、翌月の Bi-weekly① で必ずフォローし、"
        "サイクルを閉じる",
        12.5, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    # --- 運用の注記 ---
    nw = (W - GAP * 2) / 3
    for i, (lead, body) in enumerate(NOTES):
        nx = L + i * (nw + GAP)
        rect(sl, nx, NOTE_Y, nw, NOTE_H, WHITE, RULE, 0.75)
        rich(sl, nx + 0.14, NOTE_Y, nw - 0.28, NOTE_H,
             [(lead + "：", 11, True, DARK), (body, 11, False, GREY)],
             anchor=MSO_ANCHOR.MIDDLE, space=1.15)

    say(sl, L, 7.14, 7.0, 0.24, "会議体運営｜生準マネジメントサイクル", 11,
        False, GREY)
    say(sl, 12.20, 7.14, 0.72, 0.24, "2", 11, False, GREY, align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "会議体運営の1か月サイクル。第1週の Bi-weekly① で前回以降の変化点をカテゴリ別に"
        "確認し、遅れ・問題を抜き出す。第1〜2週は担当者が要因と影響範囲、リカバリ策、"
        "他案件への波及を整理する。第3週の Bi-weekly② で影響とリカバリ策を報告して方針を"
        "合意し、マネジメントへの依頼・確認事項を整理して月次報告資料をまとめる。月末の"
        "生準月次で主な変化点と対応状況を報告し、判断事項・支援依頼・前回依頼への回答を"
        "確認する。月次の決定・指摘事項は翌月の Bi-weekly① でフォローし、サイクルを閉じる。")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="meeting_cadence_daicel.pptx")
    build(ap.parse_args().output)
