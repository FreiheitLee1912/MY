# -*- coding: utf-8 -*-
"""4.1 JIRA Project Scope — 進捗・課題を1つにまとめ、証拠管理を追記した版。

Usage:
    python build_slide_jira_scope.py [-o out.pptx]

元は 01 進捗管理 / 02 課題管理 の2枚看板だったが、どちらも同じ案件に紐づく
1つの運用なので 01 に統合し、空いた 02 を証拠管理に充てる。
"""
import argparse
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
AMBER = C("FFC000")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
FAINT = C("F4F5F3")
WHITE = C("FFFFFF")

BODY, DISPLAY = "Arial", "Arial Black"
L, R = 0.42, 12.92
W = R - L

EM = 0.085            # 11pt Arial の折り返し見積り（安全側）
LH = 0.20             # 11pt / 行送り1.25 の1行


def style(run, size, bold=False, color=INK, face=None, spacing=None):
    face = face or BODY
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space=None, spacing=None, face=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face, spacing=spacing)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chip(sl, x, y, w, h, text, fill, size=11, fg=WHITE):
    sh = rect(sl, x, y, w, h, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    style(r, size, True, fg)
    return sh


def lines(text, inner_w):
    return max(1, math.ceil(len(text) / max(6, int(inner_w / EM))))


CARDS = [
    dict(no="01", title="PROGRESS & ISSUE MANAGEMENT",
         desc="Schedule commitments and the exceptions raised against them.",
         rows=[
             ("UNIT OF CONTROL",
              "One issue per milestone or task; one per exception"),
             ("WHAT IS RECORDED",
              "Baseline and actual dates, status and slippage; impact, "
              "countermeasure, owner and due date"),
             ("UPDATE RHYTHM",
              "Monthly by the Category owner; on each status change by the "
              "site"),
             ("ESCALATION PATH",
              "Slippage logged with cause and countermeasure, then routed for "
              "the decision it needs"),
             ("STATUS MODEL",
              "Not started / In progress / Done, on one scale across sites"),
         ]),
    dict(no="02", title="EVIDENCE MANAGEMENT",
         desc="The documents and decisions that justify each case.",
         rows=[
             ("UNIT OF CONTROL",
              "One evidence set per issue; no parallel folder to maintain"),
             ("WHAT IS RECORDED",
              "Approved deliverables, minutes and decision records, with "
              "version and approver"),
             ("UPDATE TRIGGER",
              "At approval or revision, by the function that owns it"),
             ("WHAT IT PROVES",
              "Who approved what, when, and on what basis"),
             ("RETENTION",
              "Held against the issue, so the audit trail stays with the case"),
         ]),
]

OUTCOMES = [
    ("Plan visibility", "Baseline, progress and actuals, readable by task."),
    ("Early issue detection", "Slippage and impact surface where work is "
                              "recorded."),
    ("Decision traceability", "Every decision carries its owner, date and "
                              "basis."),
]

# --- 版面 ---
CY, CH = 1.56, 3.76             # カードの天地
CW = (W - 0.24) / 2
CX = [L, L + CW + 0.24]
HDR_H = 0.76
LBL_W = 1.78

OUT_LBL_Y = 5.46
TILE_Y, TILE_H = 5.72, 0.76
SW_Y, SW_H = 6.62, 0.50
FOOT_Y = 7.20


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # --- ヘッダー ---
    say(sl, L, 0.36, 10.0, 0.26, "4.1 JIRA PROJECT SCOPE", 11.5, True, BLUE,
        spacing=1.0)
    chip(sl, 12.18, 0.33, 0.74, 0.28, "完", AMBER, 11, INK)
    say(sl, L, 0.62, W, 0.56,
        "One JIRA issue carries the plan, the problem and the proof",
        20, True, INK, anchor=MSO_ANCHOR.MIDDLE, face=DISPLAY)
    say(sl, L, 1.20, 12.3, 0.26,
        "HQ moves from Excel-based tracking to JIRA by October; every site "
        "records against the same fields.",
        12.5, False, GREY)

    # --- 管理対象 ---
    for card, cx in zip(CARDS, CX):
        rect(sl, cx, CY, CW, CH, WHITE, RULE, 0.75)
        rect(sl, cx, CY, CW, 0.06, BLUE)
        rect(sl, cx, CY + 0.06, CW, HDR_H - 0.06, PALE)
        say(sl, cx + 0.20, CY + 0.12, 0.60, 0.32, card["no"], 21, True, BLUE,
            face=DISPLAY)
        say(sl, cx + 0.84, CY + 0.16, CW - 1.04, 0.26, card["title"], 13.5,
            True, INK, spacing=0.4)
        say(sl, cx + 0.84, CY + 0.46, CW - 1.04, 0.24, card["desc"], 11, False,
            GREY)

        y = CY + HDR_H + 0.14
        vw = CW - 0.40 - LBL_W
        for i, (label, value) in enumerate(card["rows"]):
            h = lines(value, vw) * LH
            if i:
                rect(sl, cx + 0.20, y - 0.09, CW - 0.40, 0.01, RULE)
            say(sl, cx + 0.20, y, LBL_W, 0.22, label, 11, True, DARK,
                spacing=0.5, space=1.25)
            say(sl, cx + 0.20 + LBL_W, y, vw, h, value, 11, False, INK,
                space=1.25)
            y += max(h, 0.22) + 0.16

    # --- JIRA で何が取れるか ---
    say(sl, L, OUT_LBL_Y, 4.0, 0.22, "JIRA OUTCOMES", 11.5, True, DARK,
        spacing=1.0)
    rect(sl, L, OUT_LBL_Y + 0.24, W, 0.02, BLUE)
    tw = (W - 0.32) / 3
    for i, (head, body) in enumerate(OUTCOMES):
        tx = L + i * (tw + 0.16)
        rect(sl, tx, TILE_Y, tw, TILE_H, PALE)
        rect(sl, tx, TILE_Y, 0.07, TILE_H, BLUE)
        say(sl, tx + 0.24, TILE_Y + 0.14, tw - 0.44, 0.24, head, 12.5, True,
            DARK)
        say(sl, tx + 0.24, TILE_Y + 0.42, tw - 0.44, 0.28, body, 11, False,
            GREY, space=1.2)

    # --- So What ---
    rect(sl, L, SW_Y, W, SW_H, BLUE)
    chip(sl, L + 0.16, SW_Y + 0.07, 1.22, SW_H - 0.14, "So What", DARK, 11.5)
    say(sl, L + 1.58, SW_Y, W - 1.78, SW_H,
        "Progress, issues and evidence sit on one issue, so status and history "
        "read the same at every site",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, FOOT_Y, 7.5, 0.24, "4.1 JIRA project scope", 11, False, GREY)
    say(sl, 12.20, FOOT_Y, 0.72, 0.24, "7", 11, False, GREY,
        align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "JIRA project scope. Progress management and issue management were "
        "previously two separate blocks, but both hang off the same case, so "
        "they are merged into one: a JIRA issue per key milestone or task and "
        "one per problem, carrying planned and actual dates, progress and delay "
        "status, and for problems the impact, response plan, owner and due "
        "date. The Category owner updates monthly; the originating site updates "
        "on any status change. When something is delayed, the cause and "
        "recovery action are logged and the cross-functional input or decision "
        "needed is raised. The second block is evidence management, added here "
        "rather than left as an outcome: one evidence set per issue holding "
        "approved documents, minutes and decision records with version and "
        "approver, updated on approval or revision by the owning function, so "
        "the record of who approved what, when and on what basis travels with "
        "the case. What JIRA then gives: plan visibility, early issue "
        "detection, and decision traceability.")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="jira_scope_en.pptx")
    build(ap.parse_args().output)
