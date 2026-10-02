# -*- coding: utf-8 -*-
"""現状と課題｜需給バランス — As is / To be を1つの流れで比較する。

Usage:
    python build_slide_balance_asis_tobe.py [-o out.pptx]

As is と To be の違いは「3者が共有している文書」だけなので、
マーケ → 生準 → 拠点 の流れは1列だけ描き、その下に As is の共有文書、
To be の共有文書を重ねて差分を1か所で読ませる。
"""
import argparse
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
GREEN = C("00843D")
GREEN_PALE = C("E7F4EC")
AMBER = C("C55A11")
AMBER_FILL = C("FDF1E3")
RED = C("C00000")
RED_FILL = C("FBEAEA")
INK = C("1F2937")
GREY = C("6B7280")
SLATE = C("7F7F7F")
FAINT = C("F4F5F3")
RULE = C("D0D5DD")
WHITE = C("FFFFFF")

JP = "Meiryo"
L, R = 0.42, 12.92
W = R - L

EM11 = 0.153          # 11pt Meiryo の全角1文字幅
EM_EST = 0.158        # 折り返し見積り用（安全側）
LH11 = 0.20           # 11pt / 行送り1.25 の1行


def style(run, size, bold=False, color=INK, face=JP, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, spacing=spacing)
    return tf


def rich(sl, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=1.25):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = space
    for text, size, bold, color in parts:
        r = p.add_run(); r.text = text
        style(r, size, bold, color)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chip(sl, x, y, w, h, text, fill, size=11, fg=WHITE, outline=None, ow=1.0,
         bold=True):
    sh = rect(sl, x, y, w, h, fill, outline, ow)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    style(r, size, bold, fg)
    return sh


def bullet(sl, x, y, w, text, marker, color=INK):
    cpl = max(6, int((w - 0.20) / EM_EST))
    h = max(1, math.ceil(len(text) / cpl)) * LH11
    rect(sl, x, y + 0.075, 0.07, 0.07, marker)
    say(sl, x + 0.20, y, w - 0.20, h, text, 11, False, color, space=1.25)
    return h


NODES = [
    ("マーケ", "Marketing", "HQ営業が販売計画を日々更新"),
    ("生準", "Production Prep.", "需給を見て供給可否を判断"),
    ("拠点", "Site", "設備・投資の検討に落とす"),
]

# (チップ, 塗り, 文字色, 枠) の並び＋注記
ASIS = [
    ([("バランス（予算）", GREEN_PALE, GREEN, GREEN),
      ("バランス（月次）", AMBER_FILL, AMBER, AMBER)], None),
    ([("バランス（予算）", GREEN_PALE, GREEN, GREEN),
      ("バランス（月次）", AMBER_FILL, AMBER, AMBER)], None),
    ([("予算販売計画（予算レンジのみ）", RED_FILL, RED, RED)],
     "生準との連携がなく、レンジ外の将来案件は見えない"),
]
TOBE = [
    ([("バランス（予算）", GREEN_PALE, GREEN, GREEN),
      ("バランス（日次）", PALE, DARK, BLUE)], None),
    ([("バランス（予算）", GREEN_PALE, GREEN, GREEN),
      ("バランス（日次）", PALE, DARK, BLUE)], None),
    ([("バランス（予算）", GREEN_PALE, GREEN, GREEN),
      ("バランス（日次）", PALE, DARK, BLUE)], "3者が同じものを同じ鮮度で見る"),
]

ISSUES = ["生準が販売計画の変化を把握するまでに時間差がある",
          "拠点はレンジ内の案件しか見えず、投資の検討に織り込めない"]
GAINS = ["HQと拠点の両方が、予算レンジ外を含む将来需要を確認できる",
         "大型案件や数量変動を早期に把握し、HQ・拠点で連携して供給対応を検討できる"]

GUT = 1.38                      # 行ラベルの左ガター
CX0 = L + GUT + 0.18
COLW = 3.30
CGAP = (R - CX0 - COLW * 3) / 2

BLK_Y, BLK_H = 1.52, 3.82
NODE_Y, NODE_H = 1.76, 0.66
AS_Y, BAND_H = 2.70, 1.20
TO_Y = 4.00
EXP_Y, EXP_H = 5.46, 1.10
SW_Y, SW_H = 6.68, 0.48


def band(sl, y, label, caption, label_fill, fill, rows):
    rect(sl, L + 0.10, y, W - 0.20, BAND_H, fill)
    chip(sl, L + 0.22, y + 0.16, GUT - 0.20, 0.34, label, label_fill, 13)
    say(sl, L + 0.22, y + 0.56, GUT - 0.20, 0.40, caption, 11, False, GREY,
        align=PP_ALIGN.CENTER, space=1.15)
    for i, (chips, note) in enumerate(rows):
        cx = CX0 + i * (COLW + CGAP)
        cw = 2.96 if len(chips) == 1 else 2.40
        cy = y + 0.16
        for text, cf, fg, oc in chips:
            chip(sl, cx + (COLW - cw) / 2, cy, cw, 0.30, text, cf, 11, fg,
                 outline=oc, ow=1.0)
            cy += 0.38
        if note:
            say(sl, cx + 0.10, cy - 0.04, COLW - 0.20, 0.40, note, 11, False,
                GREY, align=PP_ALIGN.CENTER, space=1.15)


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # --- ヘッダー ---
    say(sl, L, 0.36, 9.5, 0.26, "① 現状と課題 ｜ 需給バランス", 11.5, True, BLUE,
        spacing=1.0)
    say(sl, L, 0.62, W, 0.56,
        "違いは「3者が共有する文書」だけ — 月次を日次に変え、拠点にも同じものを配る",
        20, True, INK, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L, 1.20, 12.2, 0.26,
        "販売計画はHQ営業が日々更新しているが、生準・拠点が変化点を確認するまでに"
        "時間差がある。流れそのものは変わらない。",
        12.5, False, GREY)

    # --- 比較ブロック ---
    rect(sl, L, BLK_Y, W, BLK_H, WHITE, RULE, 0.75)

    # ノード行（As is / To be 共通）
    say(sl, L + 0.22, NODE_Y + 0.12, GUT - 0.20, 0.42, "情報の流れ\n（共通）", 11,
        True, DARK, align=PP_ALIGN.CENTER, space=1.15)
    for i, (jp, en, role) in enumerate(NODES):
        cx = CX0 + i * (COLW + CGAP)
        rect(sl, cx, NODE_Y, COLW, NODE_H, DARK)
        say(sl, cx, NODE_Y + 0.07, COLW, 0.26, jp, 15, True, WHITE,
            align=PP_ALIGN.CENTER)
        say(sl, cx, NODE_Y + 0.36, COLW, 0.22, en, 11, False, PALE,
            align=PP_ALIGN.CENTER)
        say(sl, cx, NODE_Y + NODE_H + 0.06, COLW, 0.20, role, 11, False, GREY,
            align=PP_ALIGN.CENTER)
        if i < len(NODES) - 1:
            ax = cx + COLW + 0.06
            rect(sl, ax, NODE_Y + 0.20, CGAP - 0.12, 0.30, BLUE,
                 shape=MSO_SHAPE.RIGHT_ARROW)

    band(sl, AS_Y, "As is", "いま共有して\nいる文書", SLATE, FAINT, ASIS)
    band(sl, TO_Y, "To be", "これから共有\nする文書", DARK, PALE, TOBE)

    # --- 説明 ---
    bw = (W - 0.24) / 2
    for i, (title, items, accent, fill) in enumerate(
            [("As is の課題", ISSUES, RED, RED_FILL),
             ("To be でできること", GAINS, DARK, PALE)]):
        bx = L + i * (bw + 0.24)
        rect(sl, bx, EXP_Y, bw, EXP_H, WHITE, RULE, 0.75)
        rect(sl, bx, EXP_Y, 0.07, EXP_H, accent)
        say(sl, bx + 0.22, EXP_Y + 0.10, bw - 0.40, 0.22, title, 11.5, True,
            accent)
        iy = EXP_Y + 0.36
        for t in items:
            iy += bullet(sl, bx + 0.22, iy, bw - 0.40, t, accent) + 0.06

    # --- So What ---
    rect(sl, L, SW_Y, W, SW_H, BLUE)
    chip(sl, L + 0.16, SW_Y + 0.07, 1.22, SW_H - 0.14, "So What", DARK, 11.5)
    say(sl, L + 1.58, SW_Y, W - 1.78, SW_H,
        "仕組みを作り替える話ではなく、配る文書を「日次バランス」に揃えるだけで"
        "時間差と拠点の情報格差は解消する",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, 7.24, 7.0, 0.24, "現状と課題｜需給バランス", 11, False, GREY)
    say(sl, 12.20, 7.24, 0.72, 0.24, "12", 11, False, GREY, align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "需給バランスの As is / To be。マーケ→生準→拠点という情報の流れ自体は変わらず、"
        "差分は3者が共有している文書だけなので、流れは1列だけ描いて下に As is と To be の"
        "共有文書を重ねている。As is ではマーケと生準が月次バランスを共有し、拠点は予算"
        "レンジ内の予算販売計画しか持たない。そのため生準が販売計画の変化を把握するまでに"
        "時間差が生じ、拠点はレンジ外の将来案件を投資検討に織り込めない。To be は3者が同じ"
        "日次バランスを同じ鮮度で見る状態。HQと拠点の双方が予算レンジ外を含む将来需要を"
        "確認でき、大型案件や数量変動を早期に把握して供給対応を検討できる。")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="balance_asis_tobe_daicel.pptx")
    build(ap.parse_args().output)
