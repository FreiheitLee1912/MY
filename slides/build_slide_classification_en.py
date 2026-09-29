# -*- coding: utf-8 -*-
"""Classification of production-preparation cases — English, ACN layout.

Usage:
    python build_slide_classification_en.py --brand daicel [-o out.pptx]
    python build_slide_classification_en.py --brand acn    [-o out.pptx]

版面は Accenture 2020 のグリッド（左右マージン0.42"／テキスト幅12.50"、
罫線なしの Arial Black 見出し、So What バンド）。配色のみ --brand で切替。
"""
import argparse

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLACK, WHITE = C("000000"), C("FFFFFF")
INK = C("333333")

BRANDS = {
    "acn": dict(
        key=C("A100FF"), dark=C("460073"), tint=C("F4E9FF"), muted=C("70706A"),
        post=C("7500C0"), post_tint=C("EDE0FA"),
        grey=C("96968C"), grey_tint=C("E6E6DC"),
        sowhat=C("A100FF"),
    ),
    # DAICEL: brand blue from the template logo; the source slide's green is kept
    # for the post-SOP phases so its colour semantics survive the reskin.
    "daicel": dict(
        key=C("0096D8"), dark=C("00558C"), tint=C("E5F4FC"), muted=C("6B7280"),
        post=C("4C9A2A"), post_tint=C("E8F3E4"),
        grey=C("767676"), grey_tint=C("E7E7E7"),
        sowhat=C("0096D8"),
    ),
}

LAT, DISPLAY = "Arial", "Arial Black"
L, R = 0.42, 12.92
W = R - L


# ---- helpers ----------------------------------------------------------------
def style(run, size, bold=False, color=BLACK, face=LAT):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def textbox(sl, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tb, tf


def say(sl, x, y, w, h, text, size, bold=False, color=BLACK,
        align=PP_ALIGN.LEFT, face=LAT, anchor=MSO_ANCHOR.TOP, space=None):
    _, tf = textbox(sl, x, y, w, h, align, anchor)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face)
    return tf


def shape(sl, kind, x, y, w, h, fill=None, outline=None, ow=1.0):
    sh = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return sh


def box(sl, x, y, w, h, text="", size=11, bold=False, fg=BLACK, fill=WHITE,
        outline=None, align=PP_ALIGN.CENTER, space=None, ow=1.0):
    sh = shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, outline=outline, ow=ow)
    for i, ln in enumerate(str(text).split("\n")):
        p = sh.text_frame.paragraphs[0] if i == 0 else sh.text_frame.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        if ln:
            r = p.add_run(); r.text = ln
            style(r, size, bold, fg)
    return sh


# ---- the slide --------------------------------------------------------------
CASES = [
    ("1", "Potential Cases", False,
     "Manages customer themes before award. Organises needs, target products "
     "and expected SOP, then decides whether the case moves to New Launch."),
    ("2", "New Launch", True,
     "Manages mass-production preparation once the award is confirmed — spec "
     "freeze, investment, design and evaluation, PPAP and customer approval "
     "through to SOP."),
    ("3", "PCR", True,
     "Manages changes to process, equipment, production site or parts for "
     "products in mass production, from proposal and approval through to "
     "closure."),
    ("4", "Service Parts", False,
     "Manages service parts whose supply obligation continues after the model "
     "ends — demand outlook, inventory, capacity and last-production timing."),
]

PHASES = [("Inquiry", "pre"), ("Proposal\n/ Bid", "pre"), ("Investment", "pre"),
          ("Mass-Prod\nPrep", "pre"), ("SOP", "pre"), ("Mass\nProduction", "post"),
          ("Service\nParts", "out"), ("EOP", "out")]


def build(brand, output):
    P = BRANDS[brand]
    KEY, DARK, TINT, MUTED = P["key"], P["dark"], P["tint"], P["muted"]
    POST, POST_TINT = P["post"], P["post_tint"]
    GREY, GREY_TINT = P["grey"], P["grey_tint"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    say(sl, L, 0.36, 9.5, 0.24,
        "Case Classification | Scope of JIRA registration", 11.5, True, KEY)
    say(sl, L, 0.62, W, 0.66,
        "JIRA registration is limited to New Launch and PCR",
        22, True, BLACK, face=DISPLAY)
    say(sl, L, 1.34, 12.0, 0.28,
        "Production preparation manages four case types. Two of them are "
        "registered in JIRA; the other two are managed elsewhere.",
        13, False, MUTED)

    # --- four case cards ---
    say(sl, L, 1.78, 4.0, 0.26, "Case types", 12, True, KEY)
    cw = (W - 3 * 0.22) / 4
    cy = 2.10
    for i, (no, name, in_scope, body) in enumerate(CASES):
        x = L + i * (cw + 0.22)
        hd_fill = KEY if in_scope else GREY
        # header bar drawn empty: the title is set flush left so the JIRA badge
        # can sit at the right without colliding with a centred title
        box(sl, x, cy, cw, 0.44, "", fill=hd_fill, outline=None)
        say(sl, x + 0.18, cy + 0.10, cw - 1.10, 0.26, f"{no}  {name}", 13, True,
            WHITE, anchor=MSO_ANCHOR.MIDDLE)
        box(sl, x, cy + 0.44, cw, 1.36, "", fill=WHITE,
            outline=KEY if in_scope else GREY_TINT, ow=1.25 if in_scope else 1.0)
        say(sl, x + 0.16, cy + 0.56, cw - 0.32, 1.12, body, 11,
            False, INK if in_scope else MUTED, space=1.28)
        if in_scope:
            box(sl, x + cw - 0.86, cy + 0.10, 0.74, 0.24, "JIRA", 11, True,
                WHITE, DARK, None)

    # --- product lifecycle ---
    py = 4.12
    say(sl, L, py, 5.0, 0.28, "PROCESS | Product lifecycle", 13, True, DARK,
        face=DISPLAY)

    n = len(PHASES)
    pitch = W / n
    chw = pitch - 0.05
    bx = lambda i: L + i * pitch

    # milestone markers above the flow
    for i, label in [(2, "Case won"), (6, "Model end")]:
        say(sl, bx(i) - 0.95, py + 0.40, 1.90, 0.24, label, 11, True, DARK,
            align=PP_ALIGN.CENTER)
        shape(sl, MSO_SHAPE.DIAMOND, bx(i) - 0.06, py + 0.70, 0.12, 0.12, fill=DARK)

    fy = py + 0.90
    for i, (name, kind) in enumerate(PHASES):
        if kind == "pre":
            fill, out, fg = TINT, KEY, DARK
        elif kind == "post":
            fill, out, fg = POST_TINT, POST, POST
        else:
            fill, out, fg = GREY_TINT, None, GREY
        ch = shape(sl, MSO_SHAPE.CHEVRON, bx(i), fy, chw, 0.60,
                   fill=fill, outline=out, ow=1.0)
        for j, ln in enumerate(name.split("\n")):
            p = ch.text_frame.paragraphs[0] if j == 0 else ch.text_frame.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 0.92
            r = p.add_run(); r.text = ln
            style(r, 11, True, fg)

    # legend
    ly = fy + 0.76
    for i, (col, tint, txt) in enumerate(
            [(KEY, TINT, "Pre-SOP | New Launch, PPAP"),
             (POST, POST_TINT, "Post-SOP | PCR, Service Parts")]):
        x = L + i * 4.20
        shape(sl, MSO_SHAPE.RECTANGLE, x, ly + 0.03, 0.28, 0.18,
              fill=tint, outline=col, ow=1.0)
        say(sl, x + 0.40, ly, 3.90, 0.24, txt, 11, False, MUTED)

    # --- So What ---
    shape(sl, MSO_SHAPE.RECTANGLE, L, 6.22, W, 0.64, fill=P["sowhat"])
    say(sl, L + 0.30, 6.35, 1.6, 0.28, "So What", 13.5, True, WHITE,
        face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L + 2.05, 6.35, 10.3, 0.28,
        "Keep JIRA to New Launch and PCR; Potential Cases and Service Parts are "
        "managed outside it",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, 7.08, 7.0, 0.24,
        "Case Classification | Scope of JIRA registration", 11, False, MUTED)

    sl.notes_slide.notes_text_frame.text = (
        "Production preparation manages four case types. New Launch and PCR are "
        "registered in JIRA. Potential Cases are monitored separately; Service "
        "Parts are handled in their own project. Pre-SOP covers New Launch and "
        "PPAP; post-SOP covers PCR and Service Parts.")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--brand", choices=sorted(BRANDS), default="daicel")
    ap.add_argument("-o", "--output")
    a = ap.parse_args()
    build(a.brand, a.output or f"case_classification_en_{a.brand}.pptx")
