# -*- coding: utf-8 -*-
"""JIRA implementation workstreams — English, ACN style, two workstreams.

Usage:
    python build_slide_workstreams_en.py --icons <dir> [-o out.pptx]

Rebuild of the three-column original with workstream 03 (Registration and
integration rules) removed, so the two remaining workstreams get a wider column
each. Purple rules, navy headline and the dark closing band follow the source.
"""
import argparse
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


PURPLE = C("A100FF")
NAVY = C("101828")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
WHITE = C("FFFFFF")

LAT = "Arial"
L, R = 0.67, 12.67
W = R - L


def style(run, size, bold=False, color=INK, face=LAT, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, spacing=spacing)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0):
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


WORKSTREAMS = [
    ("01", "Project visibility", "visibility.png",
     "Consolidate project status and progress in reports and dashboards.",
     "Identify delays and issues early to accelerate decisions and responses."),
    ("02", "Evidence and document controls", "evidence.png",
     "Set rules for storing evidence and data, including naming, location and "
     "access rights.",
     "Provide faster access to required information and streamline audits and "
     "handovers."),
]


def build(icon_dir, output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # top rule
    rect(sl, 0, 0, 13.333, 0.05, PURPLE)

    # --- header ---
    say(sl, L + 0.25, 0.34, 7.0, 0.24, "OBJECTIVE: JIRA IMPLEMENTATION",
        11, True, PURPLE, spacing=1.2)
    rect(sl, L, 0.62, 0.07, 0.46, PURPLE)
    # 21pt keeps the headline on one line inside the 9.3" left column; the
    # milestone box starts at 10.32 so it cannot run wider than that
    say(sl, L + 0.25, 0.62, 9.30, 0.46,
        "Transition from Excel-based tracking to JIRA by October",
        21, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L + 0.25, 1.18, 9.30, 0.28,
        "Standardize project visibility and evidence management",
        13, False, GREY)

    # target milestone box
    bx, bw = 10.32, 2.35
    rect(sl, bx, 0.32, bw, 0.84, None, RULE, 1.0)
    rect(sl, bx, 0.32, 0.06, 0.84, PURPLE)
    say(sl, bx + 0.14, 0.44, bw - 0.28, 0.22, "TARGET MILESTONE",
        10.5, True, PURPLE, align=PP_ALIGN.CENTER, spacing=1.0)
    say(sl, bx + 0.14, 0.70, bw - 0.28, 0.40,
        "GLOBAL PRODUCTION\nREADINESS MEETING",
        11.5, True, NAVY, align=PP_ALIGN.CENTER, space=1.15)

    rect(sl, L, 1.62, W, 0.012, C("B9C0CC"))

    # --- section label ---
    say(sl, L, 1.82, 8.0, 0.30, "Two workstreams for JIRA implementation",
        15, True, NAVY)
    rect(sl, L, 2.16, W, 0.025, PURPLE)

    # --- two columns ---
    cw = 5.70
    gap = 0.60
    sep_x = L + cw + gap / 2
    rect(sl, sep_x, 2.44, 0.008, 3.62, RULE)

    for i, (no, title, icon, impl, outcome) in enumerate(WORKSTREAMS):
        x = L + i * (cw + gap)
        say(sl, x, 2.40, 0.95, 0.46, no, 28, True, PURPLE)
        say(sl, x + 1.00, 2.46, cw - 1.00, 0.36, title, 16, True, NAVY,
            anchor=MSO_ANCHOR.MIDDLE)
        rect(sl, x, 2.98, cw - 0.30, 0.022, PURPLE)

        say(sl, x, 3.24, 3.0, 0.22, "IMPLEMENTATION", 11, True, PURPLE, spacing=1.0)
        ip = os.path.join(icon_dir, icon)
        if os.path.exists(ip):
            sl.shapes.add_picture(ip, Inches(x + 0.06), Inches(3.56),
                                  Inches(0.62), Inches(0.62))
        say(sl, x + 0.92, 3.52, cw - 1.10, 0.80, impl, 12.5, False, INK, space=1.3)

        rect(sl, x, 4.62, cw - 0.30, 0.008, RULE)

        say(sl, x, 4.84, 3.0, 0.22, "EXPECTED OUTCOME", 11, True, PURPLE, spacing=1.0)
        say(sl, x, 5.14, cw - 0.30, 0.80, outcome, 12.5, False, INK, space=1.3)

    # --- closing band ---
    rect(sl, 0.57, 6.34, 0.09, 0.70, PURPLE)
    rect(sl, 0.66, 6.34, 12.14, 0.70, NAVY)
    say(sl, 0.90, 6.34, 11.70, 0.70,
        "Use JIRA to improve transparency, efficiency, and consistency in "
        "project delivery",
        16, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, 12.10, 7.18, 0.57, 0.22, "03", 11, True, C("9AA3AF"),
        align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "Two workstreams for the JIRA implementation. 01 Project visibility: "
        "consolidate status and progress in reports and dashboards so delays and "
        "issues surface early. 02 Evidence and document controls: set naming, "
        "location and access rules so information is found faster and audits and "
        "handovers are simpler.")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--icons", default=".", help="directory holding the icon PNGs")
    ap.add_argument("-o", "--output", default="jira_workstreams_en.pptx")
    a = ap.parse_args()
    build(a.icons, a.output)
