# -*- coding: utf-8 -*-
"""Account Management / JIRA Standard cost の日本語版（2ページ）。

Usage:
    python build_deck_account_mgmt_ja.py [-o out.pptx]

ユーザーがダイセルテンプレートで作った英語版2枚の構成と文言を日本語に起こしたもの。
P1: 導入目標時期と想定ユーザー数の回答依頼（記入表＋回答者・期限・提出先）
P2: 公式月額単価と、改定後単価での150名の計算例
"""
import argparse

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
AMBER = C("C55A11")
AMBER_FILL = C("FDF1E3")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
FAINT = C("F4F5F3")
WHITE = C("FFFFFF")

JP = "Meiryo"
L, R = 0.42, 12.92
W = R - L


def style(run, size, bold=False, color=INK, face=JP, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, spacing=spacing)
    return tf


def rich(sl, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    if space:
        p.line_spacing = space
    for text, size, bold, color in parts:
        r = p.add_run(); r.text = text
        style(r, size, bold, color)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chip(sl, x, y, w, h, text, fill, size=11, fg=WHITE, align=PP_ALIGN.CENTER):
    sh = rect(sl, x, y, w, h, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    style(r, size, True, fg)
    return sh


def is_tbd(v):
    return v.startswith("［")


def table(sl, y, cols, rows, head_h=0.38, row_h=0.42, total_last=False,
          first_bold=True, aligns=None):
    """列定義 [(見出し, 幅)] と行データで表を描く。［…］は琥珀で埋め待ちを示す。"""
    aligns = aligns or [PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * (len(cols) - 1)
    x = L
    for (name, w), al in zip(cols, aligns):
        chip(sl, x, y, w - 0.02, head_h, name, DARK, 11.5, align=al)
        x += w
    y += head_h
    for ri, row in enumerate(rows):
        total = total_last and ri == len(rows) - 1
        x = L
        for j, (val, (_, w), al) in enumerate(zip(row, cols, aligns)):
            tbd = is_tbd(val)
            fill = AMBER_FILL if tbd else (PALE if total or j == 0 else WHITE)
            rect(sl, x, y, w - 0.02, row_h, fill, RULE, 0.75)
            say(sl, x + 0.14, y, w - 0.30, row_h, val,
                13 if total and j else 12,
                (j == 0 and first_bold) or total,
                AMBER if tbd else (DARK if j == 0 or total else INK),
                align=al, anchor=MSO_ANCHOR.MIDDLE)
            x += w
        y += row_h
    return y


def header(sl, eyebrow, headline_parts, sub):
    say(sl, L, 0.36, 10.5, 0.26, eyebrow, 11.5, True, BLUE, spacing=1.0)
    rich(sl, L, 0.62, W, 0.56, headline_parts, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L, 1.20, 12.3, 0.26, sub, 12.5, False, GREY)


def foot(sl, label, page):
    say(sl, L, 7.10, 7.5, 0.24, label, 11, False, GREY)
    say(sl, 12.20, 7.10, 0.72, 0.24, page, 11, False, GREY, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- P1
ACCT_COLS = [("拠点", 2.60), ("現在の有償アカウント数", 3.10),
             ("必要アカウント数（総数）", 3.30), ("必要となる月", 3.50)]
ACCT_ROWS = [("HQ", "14", "［要確認：総数］", "［要確認：月］"),
             ("DSSE", "［要確認：現在数］", "［要確認：総数］", "［要確認：月］"),
             ("DSST", "0", "［要確認：総数］", "［要確認：月］"),
             ("DSSC", "0", "［要確認：総数］", "［要確認：月］"),
             ("DSSA", "0", "［要確認：総数］", "［要確認：月］")]
RESPONSE = [("回答者", "［拠点の担当者または役割］"),
            ("回答期限", "10月31日"),
            ("提出先", "［HQ取りまとめ担当・提出方法］")]

# ---------------------------------------------------------------- P2
RATE_COLS = [("人数帯", 3.40), ("現行単価", 3.00), ("改定後単価", 3.00),
             ("上昇率", 3.10)]
RATE_ROWS = [("1〜100", "9.05", "9.70", "+7.2%"),
             ("101〜250", "7.65", "8.20", "+7.2%"),
             ("251〜1,000", "6.40", "6.85", "+7.0%")]
CALC_COLS = [("人数帯", 3.40), ("計算", 5.00), ("月額", 4.10)]
CALC_ROWS = [("1〜100", "100名 × USD 9.70", "USD 970"),
             ("101〜150", "50名 × USD 8.20", "USD 410"),
             ("合計", "", "USD 1,380／月")]
NOTES = [
    "150名は計算方法を示す例であり、5拠点で確定した必要数ではありません。",
    "5拠点を1契約とする前提：［契約・請求単位を要確認］。別契約の場合は契約ごとに"
    "再計算します。",
    "改定単価は2026年10月13日（太平洋時間）から適用。最終費用は代理店見積・"
    "契約条件で確認します。",
]


def page1(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "アカウント管理",
           [("本日の説明を踏まえ、", 20, True, INK),
            ("導入目標時期", 20, True, BLUE),
            ("をご決定のうえ、", 20, True, INK),
            ("想定Jiraユーザー数", 20, True, BLUE),
            ("をご連絡ください", 20, True, INK)],
           "HQで5拠点の回答を集約し、展開費用を試算します。現在JIRAを利用しているのは"
           "HQとDSSEで、DSST・DSSC・DSSAへの展開を予定しています。")

    y = table(sl, 1.86, ACCT_COLS, ACCT_ROWS, head_h=0.46, row_h=0.56)

    ry = y + 0.44
    rect(sl, L, ry, W, 0.02, BLUE)
    cw = W / 3
    for i, (label, val) in enumerate(RESPONSE):
        cx = L + i * cw
        if i:
            rect(sl, cx, ry + 0.16, 0.01, 0.46, RULE)
        say(sl, cx + 0.16, ry + 0.16, 0.92, 0.46, label, 11.5, True, DARK,
            anchor=MSO_ANCHOR.MIDDLE, spacing=0.5)
        say(sl, cx + 1.08, ry + 0.16, cw - 1.20, 0.46, val, 12, True,
            AMBER if is_tbd(val) else INK, anchor=MSO_ANCHOR.MIDDLE)

    foot(sl, "アカウント管理", "16")
    sl.notes_slide.notes_text_frame.text = (
        "本日の説明を踏まえ、各拠点で導入目標時期をご決定のうえ、想定されるJiraの"
        "ユーザー数をご連絡ください。HQで5拠点の回答を集約し、展開費用を試算します。"
        "現在JIRAを利用しているのはHQとDSSEで、DSST・DSSC・DSSAへの展開を予定して"
        "います。表には、現在の有償アカウント数、必要となるアカウントの総数、その人数が"
        "必要となる月をご記入ください。回答期限は10月31日です。")


def page2(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "2. JIRA Standard の費用",
           [("月額は人数帯ごとの累進計算。改定後の単価は各人数帯で", 20, True, INK),
            ("約7%上がる", 20, True, BLUE)],
           "公式の月額単価（USD／人・月）と、改定後単価で150名を計算した場合の例。")

    rect(sl, L, 1.64, 0.10, 0.20, DARK)
    say(sl, L + 0.20, 1.60, 8.0, 0.28, "公式月額単価（USD／人・月）", 13, True,
        DARK)
    y = table(sl, 1.96, RATE_COLS, RATE_ROWS, head_h=0.36, row_h=0.38)

    rect(sl, L, y + 0.26, 0.10, 0.20, DARK)
    say(sl, L + 0.20, y + 0.22, 8.0, 0.28, "計算例：改定後単価で150名の場合", 13,
        True, DARK)
    y = table(sl, y + 0.58, CALC_COLS, CALC_ROWS, head_h=0.36, row_h=0.38,
              total_last=True)

    rect(sl, L, y + 0.14, W, 0.02, BLUE)
    ny = y + 0.24
    for n in NOTES:
        parts = []
        if "［" in n:
            a, rest = n.split("［", 1)
            b, c = rest.split("］", 1)
            parts = [(a, 11, False, GREY), ("［" + b + "］", 11, True, AMBER),
                     (c, 11, False, GREY)]
        else:
            parts = [(n, 11, False, GREY)]
        rich(sl, L + 0.14, ny, W - 0.28, 0.24, parts)
        ny += 0.27
    say(sl, L + 0.14, ny + 0.08, W - 0.28, 0.24,
        "出典：Atlassian 公式価格表および Cloud ライセンス FAQ", 11, False, GREY)

    foot(sl, "JIRA Standard の費用", "17")
    sl.notes_slide.notes_text_frame.text = (
        "JIRA Standardの公式月額単価です。1〜100名はUSD9.05から9.70、101〜250名は"
        "7.65から8.20、251〜1,000名は6.40から6.85に改定され、各人数帯で約7%の"
        "上昇です。月額は人数帯ごとの累進計算で、150名であれば100名に9.70ドル、"
        "残り50名に8.20ドルを掛けて、改定後は月額1,380ドルになります。150名は"
        "計算方法を示す例であり、5拠点で確定した必要数ではありません。5拠点を1契約と"
        "する前提ですが、契約・請求単位は要確認で、別契約の場合は契約ごとに再計算します。"
        "改定単価は2026年10月13日太平洋時間から適用され、最終費用は代理店見積と"
        "契約条件で確認します。")


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    page1(prs)
    page2(prs)
    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="account_mgmt_ja.pptx")
    build(ap.parse_args().output)
