# -*- coding: utf-8 -*-
"""As-Is — current state and issues in Global production preparation (English).

Usage:
    python build_slide_asis_en.py [-o out.pptx]

ACN版面 × DAICEL配色。左に現状（観点別）、右に課題（5項目）を置く。
"""
import argparse

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
RED = C("C00000")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
WHITE = C("FFFFFF")

LAT, DISPLAY = "Arial", "Arial Black"
L, R = 0.42, 12.92
W = R - L


def style(run, size, bold=False, color=INK, face=LAT, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        face=LAT, anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face, spacing=spacing)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0):
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def chip(sl, x, y, w, h, text, fill, size=11, fg=WHITE):
    sh = rect(sl, x, y, w, h, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    style(r, size, True, fg)
    return sh


CURRENT = [
    ("Ownership",
     "New launches and process changes are managed individually by each site."),
    ("Where information sits",
     "Case information is scattered across Excel, meeting decks and personal files."),
    ("Communication",
     "Confirmations, requests and decisions run over email; threads are not "
     "grouped by case."),
    ("Level of detail",
     "How far progress, issues and decisions are tracked differs by site."),
    ("HQ visibility",
     "HQ checks case status through meetings and individual enquiries."),
]

ISSUES = [
    ("1", "No cross-case view of progress",
     ["Current stage, open items and delays cannot be seen in one list."]),
    ("2", "Decision history and rationale are not retained",
     ["Who decided what, when and why is not organised per case."]),
    ("3", "Email dependency buries information",
     ["Key decisions and agreements are buried in threads; anyone outside the "
      "thread cannot follow the history.",
      "Repeated confirmations, forwards and CC round-trips add effort."]),
    ("4", "Level of detail and criteria differ by site",
     ["No agreed set of items to be checked globally."]),
    ("5", "Exceptions are noticed late",
     ["Undecided specs, design changes and SOP shifts are not caught early."]),
]


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # --- header ---
    say(sl, L, 0.36, 9.5, 0.26,
        "1. AS-IS | GLOBAL PRODUCTION PREPARATION", 11.5, True, BLUE, spacing=1.0)
    say(sl, L, 0.62, W, 0.62,
        "Site-by-site management over email leaves HQ without a cross-case view",
        20, True, INK, face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L, 1.28, 12.0, 0.28,
        "Progress, decision history and rationale are not retained per case.",
        13, False, GREY)

    lw, gap = 6.30, 0.30
    rx = L + lw + gap
    rw = W - lw - gap

    # --- left: current state ---
    ty = 1.78
    rect(sl, L, ty, lw, 0.34, DARK)
    say(sl, L + 0.16, ty, lw - 0.32, 0.34, "Current state", 12.5, True, WHITE,
        anchor=MSO_ANCHOR.MIDDLE)
    hy = ty + 0.34
    rect(sl, L, hy, 1.75, 0.28, PALE, RULE, 0.75)
    rect(sl, L + 1.75, hy, lw - 1.75, 0.28, PALE, RULE, 0.75)
    say(sl, L + 0.12, hy, 1.60, 0.28, "Aspect", 11, True, DARK,
        anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L + 1.87, hy, lw - 1.99, 0.28, "How it works today", 11, True, DARK,
        anchor=MSO_ANCHOR.MIDDLE)
    ry = hy + 0.28
    for aspect, state in CURRENT:
        rect(sl, L, ry, 1.75, 0.62, WHITE, RULE, 0.75)
        rect(sl, L + 1.75, ry, lw - 1.75, 0.62, WHITE, RULE, 0.75)
        say(sl, L + 0.12, ry, 1.60, 0.62, aspect, 11, True, DARK,
            anchor=MSO_ANCHOR.MIDDLE, space=1.2)
        say(sl, L + 1.87, ry, lw - 1.99, 0.62, state, 11, False, INK,
            anchor=MSO_ANCHOR.MIDDLE, space=1.25)
        ry += 0.62

    # --- right: issues ---
    rect(sl, rx, ty, rw, 0.34, RED)
    say(sl, rx + 0.16, ty, rw - 0.32, 0.34, "Issues", 12.5, True, WHITE,
        anchor=MSO_ANCHOR.MIDDLE)
    iy = ty + 0.46
    for no, title, subs in ISSUES:
        chip(sl, rx, iy, 0.30, 0.26, no, DARK, 11)
        say(sl, rx + 0.42, iy, rw - 0.42, 0.26, title, 11.5, True, DARK,
            anchor=MSO_ANCHOR.MIDDLE)
        sy = iy + 0.30
        for sub in subs:
            rect(sl, rx + 0.46, sy + 0.10, 0.10, 0.02, BLUE)
            # 11pt / 行送り1.25 で 1行 ≒ 0.24"、本文幅は約71字で折り返す
            h = 0.40 if len(sub) > 70 else 0.24
            say(sl, rx + 0.66, sy, rw - 0.66, h, sub, 11, False, INK, space=1.25)
            sy += h + 0.05
        iy = sy + 0.10

    # --- So What ---
    rect(sl, L, 6.22, W, 0.64, BLUE)
    say(sl, L + 0.30, 6.22, 1.7, 0.64, "So What", 13.5, True, WHITE,
        face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L + 2.10, 6.22, 10.3, 0.64,
        "Because the record sits in scattered email threads, problems surface "
        "only when someone asks",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, 7.08, 7.0, 0.24,
        "As-Is | Global production preparation", 11, False, GREY)
    say(sl, 12.20, 7.08, 0.72, 0.24, "1", 11, False, GREY, align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "As-Is. Cases are owned by each site and handled mainly over email, so "
        "progress, decision history and rationale are not retained per case and HQ "
        "has no cross-case view. Five issues follow from that: no cross-case "
        "progress view, no decision record, information buried in email, "
        "inconsistent level of detail between sites, and exceptions noticed late.")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="asis_global_prep_en.pptx")
    build(ap.parse_args().output)
