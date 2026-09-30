# -*- coding: utf-8 -*-
"""エグゼクティブサマリー — ACN版面 × DAICEL配色、リスク件数版.

Usage:
    python build_slide_exec_summary.py [-o out.pptx]

元スライドのKPIタイルを「リスク件数」に振り直したもの。
先頭タイルが Total risk（総数）、以降が内訳。数値は [__] のまま置き換え用に残す。
"""
import argparse

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---- DAICEL palette ---------------------------------------------------------
BLUE = C("0096D8")       # brand blue, from the standard template's logo
DARK = C("00558C")
PALE = C("E5F4FC")
RED = C("C00000")
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
WHITE = C("FFFFFF")

JP, LAT, DISPLAY = "Meiryo", "Arial", "Arial Black"

# ---- ACN grid ---------------------------------------------------------------
L, R = 0.42, 12.92
W = R - L


def style(run, size, bold=False, color=INK, face=LAT, ea=JP, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag, val in (("a:ea", ea), ("a:cs", face)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", val)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        face=LAT, anchor=MSO_ANCHOR.TOP, space=None, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face, spacing=spacing)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, kind=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


# ---- content ----------------------------------------------------------------
TILES = [
    ("Total risk ｜ リスク総数", "[__] 件", "前月比 [±__] 件", True, DARK),
    ("新規立上げ",             "[__] 件", "全体の [__]%",      False, DARK),
    ("PCR",                   "[__] 件", "全体の [__]%",      False, DARK),
    ("うち期限超過",           "[__] 件", "高リスク [__] 件",   False, RED),
]

ROWS = [
    ("総括", "[当月のリスク総数と増減を1〜2文で。"
             "例：総数は前月比 [±__] 件、期限超過が [__] 件増加]"),
    ("新規立上げ", "[フェーズ移行・SOP達成・遅延案件に紐づくリスクの要点]"),
    ("PCR", "[受付・完了の増減、滞留理由、期限超過への対策の要点]"),
]


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # --- header ---
    say(sl, L, 0.36, 6.0, 0.26, "01 ｜ サマリー", 11.5, True, BLUE, spacing=0.6)
    say(sl, L, 0.62, 8.2, 0.64, "エグゼクティブサマリー", 28, True, INK,
        face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)

    # overall status pill, right-aligned on the headline row
    pw = 4.20
    rect(sl, R - pw, 0.68, pw, 0.50, DARK, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
    say(sl, R - pw, 0.68, pw, 0.50, "全体ステータス：[順調／注意／遅延]",
        12.5, True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, 1.36, 11.0, 0.28,
        "リスクは全社合計で管理し、新規立上げ・PCR別の内訳と期限超過を併せて見る。",
        13, False, GREY)

    # --- KPI tiles ---
    ty, th = 1.80, 1.62
    tw = (W - 3 * 0.22) / 4
    for i, (label, value, note, lead, col) in enumerate(TILES):
        x = L + i * (tw + 0.22)
        rect(sl, x, ty, tw, th, PALE if lead else WHITE,
             BLUE if lead else RULE, 2.0 if lead else 1.0)
        say(sl, x + 0.22, ty + 0.18, tw - 0.44, 0.26, label, 11.5, True,
            BLUE if lead else DARK)
        say(sl, x + 0.22, ty + 0.50, tw - 0.44, 0.70, value, 36, True, col,
            face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
        say(sl, x + 0.22, ty + 1.26, tw - 0.44, 0.24, note, 11, False, GREY)

    # --- commentary ---
    cy = 3.72
    say(sl, L, cy, 6.0, 0.26, "コメント", 11.5, True, BLUE, spacing=0.6)
    ry = cy + 0.32
    for i, (label, body) in enumerate(ROWS):
        rect(sl, L, ry, W, 0.012, RULE)
        say(sl, L + 0.04, ry + 0.16, 1.90, 0.28, label, 12.5, True, DARK,
            anchor=MSO_ANCHOR.MIDDLE)
        say(sl, L + 2.05, ry + 0.16, W - 2.10, 0.40, body, 12, False, INK,
            anchor=MSO_ANCHOR.MIDDLE, space=1.25)
        ry += 0.68
    rect(sl, L, ry, W, 0.012, RULE)

    # --- So What ---
    rect(sl, L, 6.22, W, 0.64, BLUE)
    say(sl, L + 0.30, 6.22, 1.7, 0.64, "So What", 13.5, True, WHITE,
        face=DISPLAY, anchor=MSO_ANCHOR.MIDDLE)
    say(sl, L + 2.10, 6.22, 10.3, 0.64,
        "［当月の結論を1文で。例：期限超過 [__] 件への対応を優先し、"
        "リスク総数を前月水準まで戻す］",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, 7.08, 7.0, 0.24, "エグゼクティブサマリー ｜ リスク件数", 11, False, GREY)
    say(sl, 12.20, 7.08, 0.72, 0.24, "1", 11, False, GREY, align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "KPIタイルはリスク件数。先頭が Total risk（全社のリスク総数）、"
        "以降が新規立上げ・PCRの内訳と、期限超過（内数）。"
        "数値は [__] のままなので、確定値に置き換えて使う。")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="exec_summary_risk_daicel.pptx")
    build(ap.parse_args().output)
