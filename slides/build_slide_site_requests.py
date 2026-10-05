# -*- coding: utf-8 -*-
"""拠点へのお願い事項 — ACNのOpen Items行レイアウト × ダイセル配色（英語）。

Usage:
    python build_slide_site_requests.py [-o out.pptx]

元は罫線の表だったものを、ACNのOpen Itemsと同じ「番号チップ＋帯＋右端セル」の
行に置き換える。元表で空けてあった2行は、追記枠として破線で残す。
"""
import argparse
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


def C(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


BLUE = C("0096D8")
DARK = C("00558C")
PALE = C("E5F4FC")
BAND = C("EFF6FA")            # 行の地色（ACNのグレー帯にあたる）
INK = C("1F2937")
GREY = C("6B7280")
RULE = C("D0D5DD")
WHITE = C("FFFFFF")

BODY, DISPLAY = "Arial", "Arial Black"
L, R = 0.42, 12.92
W = R - L

EM = 0.085
LH = 0.20


def style(run, size, bold=False, color=INK, face=None, spacing=None):
    face = face or BODY
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, face
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", face)


def say(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space=None, spacing=None, face=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.line_spacing = space
        r = p.add_run(); r.text = ln
        style(r, size, bold, color, face, spacing=spacing)
    return tf


def rect(sl, x, y, w, h, fill, outline=None, ow=1.0, shape=MSO_SHAPE.RECTANGLE,
         dash=False):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if outline is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = outline; sh.line.width = Pt(ow)
        if dash:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def cell(sl, x, y, w, h, text, fill, size=11, fg=WHITE, bold=True,
         align=PP_ALIGN.CENTER):
    sh = rect(sl, x, y, w, h, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    style(r, size, bold, fg)
    return sh


def lines(text, inner_w):
    return max(1, math.ceil(len(text) / max(6, int(inner_w / EM))))


ROWS = [
    ("ATP reflection request",
     "Ask Marketing to reflect potential cases in the Action Plan early",
     "Marketing"),
    ("ATP data review",
     "Confirm with Marketing that changes in the input are valid - Potential C "
     "to B, volume up or down, programmes added or removed",
     "Marketing"),
    ("Balance alignment",
     "Hold a shared view of the balance position between Marketing and "
     "Production Prep. at each site",
     "Marketing + Prod. Prep."),
]
SPARE = 2                      # 元表で空けてあった追記枠

# --- 版面 ---
NO_W, SUM_W, OWN_W = 0.62, 2.60, 2.30
DET_W = W - NO_W - SUM_W - OWN_W
NO_X = L
SUM_X = NO_X + NO_W
DET_X = SUM_X + SUM_W
OWN_X = DET_X + DET_W

HEAD_Y, HEAD_H = 1.62, 0.40
ROW_Y, GAP = 2.10, 0.06
SPARE_H = 0.60
LEG_Y = 6.12
SW_Y, SW_H = 6.52, 0.50
FOOT_Y = 7.14


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # --- ヘッダー ---
    say(sl, L, 0.36, 10.5, 0.26, "REQUESTS TO THE SITES", 11.5, True, BLUE,
        spacing=1.0)
    say(sl, L, 0.62, W, 0.56,
        "Three things we ask of each site to strengthen global supply management",
        20, True, INK, anchor=MSO_ANCHOR.MIDDLE, face=DISPLAY)
    say(sl, L, 1.20, 12.3, 0.26,
        "Each item is an action for the site to take with its Marketing "
        "counterpart; none of them needs a new system.",
        12.5, False, GREY)

    # --- 見出し行 ---
    for x, w, t, al in ((NO_X, NO_W, "No.", PP_ALIGN.CENTER),
                        (SUM_X, SUM_W, "Request", PP_ALIGN.LEFT),
                        (DET_X, DET_W, "What we ask", PP_ALIGN.LEFT),
                        (OWN_X, OWN_W, "Counterpart", PP_ALIGN.CENTER)):
        cell(sl, x, HEAD_Y, w - 0.02, HEAD_H, t, DARK, 11.5, WHITE, True, al)

    # --- 依頼事項 ---
    y = ROW_Y
    for i, (summary, detail, owner) in enumerate(ROWS, 1):
        h = max(0.76, lines(detail, DET_W - 0.36) * LH + 0.44)
        cell(sl, NO_X, y, NO_W - 0.02, h, str(i), DARK, 15, WHITE)
        rect(sl, SUM_X, y, SUM_W + DET_W - 0.02, h, BAND)
        say(sl, SUM_X + 0.16, y, SUM_W - 0.32, h, summary, 11.5, True, DARK,
            anchor=MSO_ANCHOR.MIDDLE, space=1.2)
        say(sl, DET_X + 0.04, y, DET_W - 0.36, h, detail, 11, False, INK,
            anchor=MSO_ANCHOR.MIDDLE, space=1.25)
        cell(sl, OWN_X, y, OWN_W, h, owner, BLUE, 11, WHITE)
        y += h + GAP

    for _ in range(SPARE):
        rect(sl, NO_X, y, NO_W - 0.02, SPARE_H, WHITE, RULE, 0.75, dash=True)
        rect(sl, SUM_X, y, SUM_W + DET_W - 0.02, SPARE_H, WHITE, RULE, 0.75,
             dash=True)
        rect(sl, OWN_X, y, OWN_W, SPARE_H, WHITE, RULE, 0.75, dash=True)
        say(sl, SUM_X + 0.16, y, SUM_W + DET_W - 0.32, SPARE_H,
            "Open for further items", 11, False, GREY,
            anchor=MSO_ANCHOR.MIDDLE)
        y += SPARE_H + GAP

    # --- 凡例 ---
    rect(sl, L, LEG_Y + 0.02, 0.62, 0.20, DARK)
    say(sl, L + 0.74, LEG_Y, 4.0, 0.24, "= item number on this list", 11, False,
        GREY, anchor=MSO_ANCHOR.MIDDLE)
    rect(sl, 5.20, LEG_Y + 0.02, 0.62, 0.20, BLUE)
    say(sl, 5.94, LEG_Y, 6.0, 0.24,
        "= the function the site works with on the item", 11, False, GREY,
        anchor=MSO_ANCHOR.MIDDLE)

    # --- So What ---
    rect(sl, L, SW_Y, W, SW_H, BLUE)
    cell(sl, L + 0.16, SW_Y + 0.07, 1.22, SW_H - 0.14, "So What", DARK, 11.5)
    say(sl, L + 1.58, SW_Y, W - 1.78, SW_H,
        "All three keep the ATP input current, so the balance every site "
        "works from reflects the latest demand",
        13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    say(sl, L, FOOT_Y, 7.5, 0.24, "Requests to the sites", 11, False, GREY)
    say(sl, 12.20, FOOT_Y, 0.72, 0.24, "12", 11, False, GREY,
        align=PP_ALIGN.RIGHT)

    sl.notes_slide.notes_text_frame.text = (
        "Requests to the sites, to strengthen global supply management. Three "
        "items. First, ATP reflection: ask Marketing to reflect potential cases "
        "in the Action Plan early. Second, ATP data review: confirm with "
        "Marketing that changes in the input are valid - a potential moving from "
        "C to B, volume moving up or down, programmes added or removed. Third, "
        "balance alignment: hold a shared view of the balance position between "
        "Marketing and Production Preparation at each site. Each item is an "
        "action the site takes with its Marketing counterpart and none of them "
        "needs a new system. Two rows are left open for items that come up "
        "later. Together the three keep the ATP input current, so the balance "
        "each site works to reflects the latest demand.")

    prs.save(output)
    print("saved", output)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", default="site_requests_en.pptx")
    build(ap.parse_args().output)
